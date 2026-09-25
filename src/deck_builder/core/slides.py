from pathlib import Path

from pptx.presentation import Presentation
from pptx.slide import Slide, SlideLayout
from pptx.util import Emu, Pt

from deck_builder.configs.base import PresentationConfig
from deck_builder.configs.schemas import BookInfo
from deck_builder.core.images import add_cover_with_backdrop, add_right_margin_icon
from deck_builder.core.texts import (
    add_accent_text_box,
    add_safe_textbox,
    normalize_cover_type,
    parse_reasons_and_tags,
)


def add_title_slide(
    presentation: Presentation,
    *,
    config: PresentationConfig,
    title_text: str,
    suptitle_text: str,
    attribution_text: str,
    logo_path: Path,
) -> Slide:
    # 6 — index of blank slide layout in python-pptx
    slide_layout: SlideLayout = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(slide_layout)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = config.colors.blue
    full_text_width = (
        config.slide_size.width - config.margins.left - config.margins.right
    )
    _ = slide.shapes.add_picture(
        str(logo_path),
        left=config.margins.left,
        top=config.title_slide.logo_top,
        height=config.title_slide.logo_height,
    )
    _ = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.margins.left,
        position_top=config.title_slide.title_top,
        width=Emu(full_text_width),
        text=title_text,
        font_size=config.typography.header_1,
        color=config.colors.white,
        bold=True,
    )
    _ = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.margins.left,
        position_top=config.title_slide.suptitle_top,
        width=Emu(full_text_width),
        text=suptitle_text,
        font_size=config.typography.header_2,
        color=config.colors.white,
    )
    _ = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.margins.left,
        position_top=config.title_slide.attribution_top,
        width=Emu(full_text_width),
        text=attribution_text,
        font_size=config.typography.header_3,
        color=config.colors.white,
    )
    return slide


def add_bookinfo_slide(
    presentation: Presentation,
    *,
    config: PresentationConfig,
    book: BookInfo,
) -> Slide:
    slide_layout = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(slide_layout)
    _ = add_right_margin_icon(
        slide=slide,
        config=config,
    )
    reasons, tags = parse_reasons_and_tags(book.reasons)
    # автор
    _ = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.text_box_left,
        position_top=config.book_info_slide.author_top,
        width=config.book_info_slide.text_box_width,
        text=book.authors,
        font_size=config.typography.header_4,
        color=config.colors.black,
    )
    # название книги
    _ = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.text_box_left,
        position_top=config.book_info_slide.title_top,
        width=config.book_info_slide.text_box_width,
        text=book.title,
        font_size=config.typography.header_3,
        color=config.colors.black,
        bold=True,
    )
    # аннотация
    annotation_box = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.text_box_left,
        position_top=config.book_info_slide.annotation_top,
        width=config.book_info_slide.text_box_width,
        text=book.annotation,
        font_size=config.typography.text,
        color=config.colors.black,
    )
    for paragraph in annotation_box.text_frame.paragraphs:
        paragraph.space_before = Pt(5)
    # плашка с причинами купить
    _ = add_accent_text_box(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.text_box_left,
        position_top=config.book_info_slide.reasons_accent_top,
        text_items=["5 причин купить"],
        font_size=config.typography.header_4,
        text_color=config.colors.white,
        accent_color=config.colors.blue,
    )
    reasons_text_box = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.text_box_left,
        position_top=config.book_info_slide.reasons_top,
        width=config.book_info_slide.text_box_width,
        text=reasons,
        font_size=config.typography.text,
        color=config.colors.black,
    )
    for paragraph in reasons_text_box.text_frame.paragraphs:
        paragraph.space_before = Pt(5)
    # параметры книги
    params_text = (
        f"ISBN {book.isbn} • "
        f"Серия: {book.series}\n"
        f"{book.book_format} • {int(book.pages)} с. • "
        f"{normalize_cover_type(book.binding)} • "
        f"Красочность: {book.colority} • Тираж: {book.print_run} • Возраст: {book.age_limit}"
    )
    _ = add_safe_textbox(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.text_box_left,
        position_top=config.book_info_slide.params_top,
        width=config.book_info_slide.text_box_width,
        text=params_text,
        font_size=config.typography.text,
        color=config.colors.black,
    )
    # теги
    _ = add_accent_text_box(
        slide=slide,
        config=config,
        position_left=config.book_info_slide.tags_left,
        position_top=config.book_info_slide.tags_top,
        text_items=tags,
        font_size=config.typography.header_4,
        text_color=config.colors.white,
        accent_color=config.colors.blue,
    )
    # обложка
    _ = add_cover_with_backdrop(
        slide=slide,
        config=config,
        item_code=book.item_code,
        cover_position_left=config.book_info_slide.cover_left,
        cover_position_top=config.book_info_slide.cover_top,
        cover_max_width=config.book_info_slide.cover_max_width,
        cover_max_height=config.book_info_slide.cover_max_height,
        backdrop_position_left=config.book_info_slide.backdrop_left,
        backdrop_position_top=config.book_info_slide.backdrop_top,
    )
    return slide
