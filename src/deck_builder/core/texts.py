from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.slide import Slide
from pptx.util import Emu, Length, Mm

from deck_builder.configs.base import PresentationConfig

MIN_HEIGHT: Length = Mm(5)

ACCENT_BOX_MIN_WIDTH: Length = Mm(37.5)
ACCENT_BOX_HEIGHT: Length = Mm(6)
ACCENT_BOX_GAP: Length = Mm(5)
CHAR_WIDTH_FACTOR: float = 2.5


def add_safe_textbox(
    slide: Slide,
    *,
    config: PresentationConfig,
    position_left: Length,
    position_top: Length,
    width: Length,
    text: str,
    font_size: Length,
    color: RGBColor,
    bold: bool = False,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
    min_height: Length = MIN_HEIGHT,
) -> Shape:
    textbox: Shape = slide.shapes.add_textbox(
        left=position_left,
        top=position_top,
        width=width,
        height=min_height,
    )

    text_frame = textbox.text_frame
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.alignment = alignment
    font = paragraph.font
    font.size = font_size or config.typography.text
    font.name = config.typography.font_family
    font.bold = bold
    font.color.rgb = color or config.colors.black
    return textbox


def add_accent_text_box(
    slide: Slide,
    *,
    config: PresentationConfig,
    position_left: Length,
    position_top: Length,
    text_items: list[str],
    font_size: Length,
    text_color: RGBColor,
    accent_color: RGBColor,
    bold: bool = True,
    min_width: Length = ACCENT_BOX_MIN_WIDTH,
    box_height: Length = ACCENT_BOX_HEIGHT,
    gap: Length = ACCENT_BOX_GAP,
    char_width_factor: float = CHAR_WIDTH_FACTOR,
) -> list[BaseShape]:
    shapes: list[BaseShape] = []
    current_x: Length = position_left
    for text in text_items:
        width: Length = Mm(
            max(
                len(text) * char_width_factor,
                min_width.mm,
            )
        )
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            left=current_x,
            top=position_top,
            width=width,
            height=box_height,
        )
        box.fill.solid()
        box.fill.fore_color.rgb = accent_color or config.colors.blue
        box.line.fill.background()
        box.shadow.inherit = False
        box.rotation = 359
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = text
        font = paragraph.font
        font.size = font_size
        font.name = config.typography.font_family
        font.bold = bold
        font.color.rgb = text_color
        shapes.append(box)
        current_x = Emu(current_x + width + gap)
    return shapes


def parse_reasons_and_tags(
    reasons: str,
    *,
    min_tags: int = 2,
    placeholder: str = "<не заполнен тег>",
) -> tuple[str, list[str]]:
    if not reasons:
        return "", []
    lines = reasons.splitlines()
    tags = [
        line.strip().lstrip("#").strip().replace("_", " ")
        for line in lines
        if line.strip().startswith("#")
    ]
    cleaned_reasons = "\n".join(
        line for line in lines if not line.strip().startswith("#")
    )
    if not tags:
        tags = [placeholder] * min_tags
    while len(tags) < min_tags:
        tags.append(placeholder)
    return cleaned_reasons, tags


def normalize_cover_type(cover_type: str) -> str:
    return "мягкая обложка" if cover_type == "3" else cover_type
