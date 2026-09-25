from functools import lru_cache
from http.client import HTTPResponse
from io import BytesIO
from typing import cast
from urllib.error import URLError
from urllib.request import Request, urlopen

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.shapes.groupshape import CT_GroupShape
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.picture import Picture
from pptx.slide import Slide
from pptx.util import Emu, Length, Mm

from deck_builder.configs.base import PresentationConfig


def send_to_back(
    slide: Slide,
    shape: BaseShape,
) -> None:
    # python-pptx has no public API for changing arbitrary shape z-order.
    sp_tree: CT_GroupShape = (
        slide.shapes._spTree  # pyright: ignore[reportPrivateUsage]
    )
    sp_tree.remove(  # pyright: ignore[reportUnknownMemberType]
        shape._element  # pyright: ignore[reportPrivateUsage]
    )
    sp_tree.insert(  # pyright: ignore[reportUnknownMemberType]
        2,
        shape._element,  # pyright: ignore[reportPrivateUsage]
    )


@lru_cache(maxsize=256)
def download_cover_from_cdn(
    item_code: str,
    *,
    timeout: float = 10.0,
) -> bytes | None:
    cover_url = f"https://cdn.ast.ru/v2/{item_code}/COVER/cover1.jpg"

    request = Request(
        cover_url,
        headers={
            "User-Agent": "deck-builder/1.0",
        },
    )

    try:
        # URL always uses the fixed HTTPS cdn.ast.ru host.
        with cast(
            HTTPResponse,
            urlopen(request, timeout=timeout),  # nosec B310 -- URL is restricted to HTTPS cdn.ast.ru above.
        ) as response:
            return response.read()
    except URLError, TimeoutError, ValueError:
        return None


def calculate_image_fit(
    image_data: bytes,
    *,
    max_width: Length,
    max_height: Length,
) -> tuple[Length, Length]:
    with Image.open(BytesIO(image_data)) as image:
        image_width, image_height = image.size

    image_ratio = image_width / image_height
    container_ratio = max_width / max_height

    if image_ratio >= container_ratio:
        width = max_width
        height = Emu(round(max_width / image_ratio))
    else:
        height = max_height
        width = Emu(round(max_height * image_ratio))

    return width, height


def add_right_margin_icon(
    slide: Slide,
    *,
    config: PresentationConfig,
) -> tuple[BaseShape, BaseShape]:
    right_margin: Shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=Emu(config.slide_size.width - config.margins.right),
        top=Mm(0),
        width=config.margins.right,
        height=config.slide_size.height,
    )
    right_margin.fill.solid()
    right_margin.fill.fore_color.rgb = config.colors.blue
    right_margin.shadow.inherit = False
    right_margin.line.fill.background()

    icon: Picture = slide.shapes.add_picture(
        str(config.assets.icon),
        left=config.book_info_slide.icon_left,
        top=config.book_info_slide.icon_top,
        height=config.book_info_slide.icon_height,
    )

    return right_margin, icon


def add_missing_cover_placeholder(
    slide: Slide,
    *,
    config: PresentationConfig,
    position_left: Length,
    position_top: Length,
    width: Length,
    height: Length,
    text: str = "НЕТ ИЗОБРАЖЕНИЯ",
) -> Shape:
    placeholder: Shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=position_left,
        top=position_top,
        width=width,
        height=height,
    )

    placeholder.fill.solid()
    placeholder.fill.fore_color.rgb = config.colors.beige
    placeholder.line.fill.background()
    placeholder.shadow.inherit = False

    paragraph = placeholder.text_frame.paragraphs[0]
    paragraph.text = text

    font = paragraph.font
    font.size = config.typography.header_4
    font.name = config.typography.font_family
    font.bold = True
    font.color.rgb = config.colors.black

    return placeholder


def add_cover_backdrop(
    slide: Slide,
    *,
    config: PresentationConfig,
    position_left: Length,
    position_top: Length,
    width: Length,
    height: Length,
) -> Shape:
    backdrop: Shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=position_left,
        top=position_top,
        width=width,
        height=height,
    )

    backdrop.fill.solid()
    backdrop.fill.fore_color.rgb = config.colors.blue
    backdrop.shadow.inherit = False
    backdrop.line.fill.background()

    send_to_back(slide, backdrop)

    return backdrop


def add_cover_with_backdrop(
    slide: Slide,
    *,
    config: PresentationConfig,
    item_code: str,
    cover_position_left: Length,
    cover_position_top: Length,
    cover_max_width: Length,
    cover_max_height: Length,
    backdrop_position_left: Length,
    backdrop_position_top: Length,
) -> tuple[Shape, Picture | Shape]:
    cover_data = download_cover_from_cdn(item_code)

    if cover_data is not None:
        cover_width, cover_height = calculate_image_fit(
            cover_data,
            max_width=cover_max_width,
            max_height=cover_max_height,
        )

        cover: Picture | Shape = slide.shapes.add_picture(
            BytesIO(cover_data),
            left=cover_position_left,
            top=cover_position_top,
            width=cover_width,
            height=cover_height,
        )
        cover.shadow.inherit = False

    else:
        cover_width = cover_max_width
        cover_height = cover_max_height

        cover = add_missing_cover_placeholder(
            slide,
            config=config,
            position_left=cover_position_left,
            position_top=cover_position_top,
            width=cover_width,
            height=cover_height,
        )

    backdrop = add_cover_backdrop(
        slide,
        config=config,
        position_left=backdrop_position_left,
        position_top=backdrop_position_top,
        width=cover_width,
        height=cover_height,
    )

    return backdrop, cover
