from dataclasses import dataclass
from pathlib import Path
from typing import cast

from polars import DataFrame
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType

from deck_builder.configs.base import PresentationConfig
from deck_builder.configs.schemas import BookInfo, BookRow, TitlesDigestSchema
from deck_builder.configs.titles_digest import TITLES_DIGEST_CONFIG
from deck_builder.core.slides import add_bookinfo_slide, add_title_slide
from deck_builder.io.export import to_pptx
from deck_builder.io.load import load_from_excel

DEFAULT_TITLE_TEXT: str = "ДАЙДЖЕСТ НОВИНОК"
DEFAULT_SUPTITLE_TEXT: str = "РЕДАКЦИИ «ЛИНГВА»"
DEFAULT_ATTRIBUTION_TEXT = "МЕСЯЦ ГОД"


@dataclass(frozen=True, slots=True)
class TitlesDigest:
    config: PresentationConfig
    deck: PresentationType
    df: DataFrame | None

    @classmethod
    def from_excel(
        cls,
        file_path: str,
    ) -> TitlesDigest:
        df = load_from_excel(
            file_path=str(file_path),
            schema=TitlesDigestSchema,
        )
        deck = Presentation()

        return cls(
            config=TITLES_DIGEST_CONFIG,
            deck=deck,
            df=df,
        )

    def build(
        self,
        *,
        title_text: str = DEFAULT_TITLE_TEXT,
        suptitle_text: str = DEFAULT_SUPTITLE_TEXT,
        attribution_text: str = DEFAULT_ATTRIBUTION_TEXT,
    ) -> TitlesDigest:
        self.deck.slide_height = self.config.slide_size.height
        self.deck.slide_width = self.config.slide_size.width

        _ = add_title_slide(
            logo_path=self.config.assets.logo,
            presentation=self.deck,
            config=self.config,
            title_text=title_text,
            suptitle_text=suptitle_text,
            attribution_text=attribution_text,
        )

        if self.df is not None:
            for raw_row in self.df.iter_rows(named=True):
                row = cast(BookRow, cast(object, raw_row))
                book = BookInfo(
                    item_code=row["Код ном-ры"],
                    isbn=row["ISBN"],
                    title=row["Наименование на обложку"],
                    authors=row["Авторы на обложку"],
                    annotation=row["Аннотация"],
                    reasons=row["Пять причин купить"],
                    series=row["Серия"],
                    book_format=row["Формат"],
                    binding=row["Продукция.Тип переплета"],
                    pages=row["Кол-во стр"],
                    colority=row["Красочность блока текста"],
                    print_run=row["Тираж"],
                    age_limit=row["Продукция.Возрастное ограничение"],
                )
                _ = add_bookinfo_slide(
                    presentation=self.deck,
                    config=self.config,
                    book=book,
                )
        return self

    def to_pptx(self, output_path: str | Path) -> None:
        to_pptx(self.deck, output_path)
