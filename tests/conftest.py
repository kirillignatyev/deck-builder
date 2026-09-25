from dataclasses import replace
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock
import importlib
import sys

import pytest
import polars as pl
from PIL import Image
from pptx import Presentation
from deck_builder.configs.base import PresentationAssets
from deck_builder.configs.schemas import BookInfo, TitlesDigestSchema
from deck_builder.configs.titles_digest import TITLES_DIGEST_CONFIG
from deck_builder.core import images


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    images.download_cover_from_cdn.cache_clear()
    monkeypatch.setattr(
        images,
        "urlopen",
        Mock(side_effect=AssertionError("Network forbidden in tests")),
    )
    yield
    images.download_cover_from_cdn.cache_clear()


@pytest.fixture
def config():
    assets = Path(__file__).resolve().parents[1] / "src/deck_builder/assets"
    return replace(
        TITLES_DIGEST_CONFIG,
        assets=PresentationAssets(
            assets / "logo_lingua.png", assets / "icon_lingua.png"
        ),
    )


@pytest.fixture
def deck():
    return Presentation()


@pytest.fixture
def slide(deck):
    return deck.slides.add_slide(deck.slide_layouts[6])


@pytest.fixture
def png():
    def make(size=(200, 100)):
        stream = BytesIO()
        Image.new("RGB", size, "blue").save(stream, format="PNG")
        return stream.getvalue()

    return make


@pytest.fixture
def row():
    return dict(
        zip(
            TitlesDigestSchema.dtypes,
            [
                "00123",
                "9781234567890",
                "Книга",
                "Автор",
                "Аннотация",
                "Причина\n#тег",
                "Серия",
                "60x90",
                "3",
                128,
                "1+1",
                1000,
                "12+",
            ],
            strict=True,
        )
    )


@pytest.fixture
def frame(row):
    return pl.DataFrame([row], schema=TitlesDigestSchema.dtypes)


@pytest.fixture
def book(row):
    return BookInfo(
        **dict(zip(BookInfo.__dataclass_fields__, row.values(), strict=True))
    )


@pytest.fixture
def isolated_slides():
    return importlib.import_module("deck_builder.core.slides")


@pytest.fixture
def builder(isolated_slides):
    name = "deck_builder.builders.titles_digest"
    sys.modules.pop(name, None)
    module = importlib.import_module(name)
    yield module
    sys.modules.pop(name, None)
