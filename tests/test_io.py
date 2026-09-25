from unittest.mock import Mock
import pytest
import polars as pl
from deck_builder.configs.schemas import TitlesDigestSchema
from deck_builder.io import load, validate, export


@pytest.mark.parametrize(
    "required,present,expected",
    [
        ([], [], True),
        (["a"], ["a", "b"], True),
        (["a", "b"], ["a"], False),
        (["a", "a"], ["a"], True),
    ],
)
def test_required(required, present, expected, capsys):
    assert (
        validate.has_required_columns(required, pl.DataFrame({c: [] for c in present}))
        is expected
    )
    assert bool(capsys.readouterr().out) is (not expected)


def test_valid_cast(frame):
    assert validate.has_valid_dtypes(TitlesDigestSchema, frame)
    assert validate.has_valid_dtypes(
        TitlesDigestSchema, frame.with_columns(pl.col("Кол-во стр").cast(pl.String))
    )


@pytest.mark.parametrize("value", ["bad", "32768"])
def test_invalid_cast(frame, value, capsys):
    assert not validate.has_valid_dtypes(
        TitlesDigestSchema, frame.with_columns(pl.lit(value).alias("Кол-во стр"))
    )
    assert "Кол-во стр → Int16" in capsys.readouterr().out


def test_nulls_allowed(frame):
    assert validate.has_valid_dtypes(
        TitlesDigestSchema, frame.with_columns(pl.lit(None).alias("Кол-во стр"))
    )


def test_load_casts_preserves_extra(monkeypatch, frame):
    incoming = frame.with_columns(
        pl.col("Кол-во стр").cast(pl.String), pl.lit("extra").alias("extra")
    )
    reader = Mock(return_value=incoming)
    monkeypatch.setattr(load, "read_excel", reader)
    actual = load.load_from_excel("input.xlsx", TitlesDigestSchema)
    reader.assert_called_once_with("input.xlsx")
    assert actual["Кол-во стр"].dtype == pl.Int16
    assert actual["Код ном-ры"].item() == "00123"
    assert actual["extra"].item() == "extra"
    assert incoming["Кол-во стр"].dtype == pl.String


@pytest.mark.parametrize("invalid", ["missing", "dtype"])
def test_load_invalid(monkeypatch, frame, invalid):
    incoming = (
        frame.drop("ISBN")
        if invalid == "missing"
        else frame.with_columns(pl.lit("bad").alias("Тираж"))
    )
    monkeypatch.setattr(load, "read_excel", Mock(return_value=incoming))
    assert load.load_from_excel("input.xlsx", TitlesDigestSchema) is None


@pytest.mark.parametrize(
    "error", [FileNotFoundError("missing"), ValueError("bad workbook")]
)
def test_load_errors(monkeypatch, error):
    monkeypatch.setattr(load, "read_excel", Mock(side_effect=error))
    with pytest.raises(type(error), match=str(error)):
        load.load_from_excel("input.xlsx", TitlesDigestSchema)


@pytest.mark.parametrize(
    "name,expected",
    [
        ("deck", "deck.pptx"),
        ("deck.zip", "deck.pptx"),
        ("deck.PPTX", "deck.PPTX"),
        ("deck.pptx", "deck.pptx"),
    ],
)
@pytest.mark.parametrize("as_string", [False, True])
def test_export(tmp_path, name, expected, as_string):
    deck = Mock()
    path = tmp_path / "nested" / name
    assert export.to_pptx(deck, str(path) if as_string else path) is None
    assert path.parent.is_dir()
    deck.save.assert_called_once_with(str(path.parent / expected))


def test_export_failure(tmp_path):
    with pytest.raises(OSError, match="denied"):
        export.to_pptx(
            Mock(save=Mock(side_effect=OSError("denied"))), tmp_path / "a.pptx"
        )


def test_real_pptx_roundtrip(deck, tmp_path):
    from pptx import Presentation

    deck.slides.add_slide(deck.slide_layouts[6])
    export.to_pptx(deck, tmp_path / "nested/deck")
    assert len(Presentation(tmp_path / "nested/deck.pptx").slides) == 1
