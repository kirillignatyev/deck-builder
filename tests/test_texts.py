import pytest
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Mm, Pt
from deck_builder.core.texts import (
    add_safe_textbox,
    add_accent_text_box,
    parse_reasons_and_tags,
    normalize_cover_type,
)


@pytest.mark.parametrize(
    "value,expected",
    [
        ("", ("", [])),
        ("reason", ("reason", ["<не заполнен тег>"] * 2)),
        (
            "  reason  \n #one_two\n##three\nlast",
            ("  reason  \nlast", ["one two", "three"]),
        ),
        ("#a\n#b\n#c", ("", ["a", "b", "c"])),
        ("a\r\nb\r#tag", ("a\nb", ["tag", "<не заполнен тег>"])),
        ("#", ("", ["", "<не заполнен тег>"])),
    ],
)
def test_parse(value, expected):
    assert parse_reasons_and_tags(value) == expected


@pytest.mark.parametrize("minimum", [0, 1, 3])
def test_custom_tag_padding(minimum):
    assert parse_reasons_and_tags(
        "reason", min_tags=minimum, placeholder="missing"
    ) == ("reason", ["missing"] * minimum)


@pytest.mark.parametrize(
    "value,expected",
    [
        ("3", "мягкая обложка"),
        ("7", "7"),
        ("", ""),
        (" 3", " 3"),
        ("Твердая", "Твердая"),
    ],
)
def test_binding(value, expected):
    assert normalize_cover_type(value) == expected


@pytest.mark.parametrize(
    "bold,alignment", [(False, PP_ALIGN.LEFT), (True, PP_ALIGN.CENTER)]
)
def test_textbox(slide, config, bold, alignment):
    box = add_safe_textbox(
        slide,
        config=config,
        position_left=Mm(2),
        position_top=Mm(3),
        width=Mm(40),
        text="Привет",
        font_size=Pt(14),
        color=config.colors.blue,
        bold=bold,
        alignment=alignment,
    )
    assert (box.left, box.top, box.width, box.height) == (Mm(2), Mm(3), Mm(40), Mm(5))
    assert box.text == "Привет"
    assert box.text_frame.word_wrap
    assert box.text_frame.auto_size == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    p = box.text_frame.paragraphs[0]
    assert (p.alignment, p.font.bold, p.font.size, p.font.name, p.font.color.rgb) == (
        alignment,
        bold,
        Pt(14),
        config.typography.font_family,
        config.colors.blue,
    )


def test_textbox_fallback(slide, config):
    box = add_safe_textbox(
        slide,
        config=config,
        position_left=0,
        position_top=0,
        width=Mm(40),
        text=123,
        font_size=None,
        color=None,
        min_height=Mm(8),
    )
    assert box.text == "123"
    assert box.height == Mm(8)
    assert box.text_frame.paragraphs[0].font.size == config.typography.text
    assert box.text_frame.paragraphs[0].font.color.rgb == config.colors.black


@pytest.mark.parametrize("items", [[], ["tag"], ["short", "a" * 30]])
def test_accents(slide, config, items):
    boxes = add_accent_text_box(
        slide,
        config=config,
        position_left=Mm(2),
        position_top=Mm(3),
        text_items=items,
        font_size=Pt(12),
        text_color=config.colors.white,
        accent_color=config.colors.blue,
    )
    assert len(boxes) == len(items)
    x = Mm(2)
    for box, text in zip(boxes, items):
        assert (box.text, box.left, box.top, box.width, box.height) == (
            text,
            x,
            Mm(3),
            Mm(max(len(text) * 2.5, 37.5)),
            Mm(6),
        )
        assert box.fill.fore_color.rgb == config.colors.blue
        assert box.rotation == 359
        assert box.text_frame.paragraphs[0].font.bold
        x += box.width + Mm(5)


def test_accent_overrides(slide, config):
    (box,) = add_accent_text_box(
        slide,
        config=config,
        position_left=0,
        position_top=0,
        text_items=["abc"],
        font_size=Pt(9),
        text_color=config.colors.black,
        accent_color=None,
        bold=False,
        min_width=Mm(2),
        box_height=Mm(8),
        gap=Mm(1),
        char_width_factor=4,
    )
    assert (box.width, box.height) == (Mm(12), Mm(8))
    assert box.fill.fore_color.rgb == config.colors.blue
    assert box.text_frame.paragraphs[0].font.bold is False
