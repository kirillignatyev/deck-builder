from unittest.mock import MagicMock, Mock
from urllib.error import URLError
import pytest
from PIL import UnidentifiedImageError
from pptx.util import Mm
from deck_builder.core import images


@pytest.mark.parametrize(
    "size,expected",
    [((200, 100), (100, 50)), ((100, 200), (40, 80)), ((100, 80), (100, 80))],
)
def test_fit(png, size, expected):
    assert images.calculate_image_fit(
        png(size), max_width=Mm(100), max_height=Mm(80)
    ) == tuple(Mm(x) for x in expected)


def test_invalid_image():
    with pytest.raises(UnidentifiedImageError):
        images.calculate_image_fit(b"bad", max_width=Mm(10), max_height=Mm(10))


def test_download_request_cache(monkeypatch):
    response = MagicMock()
    response.__enter__.return_value.read.return_value = b"cover"
    opener = Mock(return_value=response)
    monkeypatch.setattr(images, "urlopen", opener)
    assert images.download_cover_from_cdn("001", timeout=2) == b"cover"
    assert images.download_cover_from_cdn("001", timeout=2) == b"cover"
    opener.assert_called_once()
    req = opener.call_args.args[0]
    assert req.full_url == "https://cdn.ast.ru/v2/001/COVER/cover1.jpg"
    assert req.get_header("User-agent") == "deck-builder/1.0"
    assert opener.call_args.kwargs == {"timeout": 2}
    response.__exit__.assert_called_once()
    images.download_cover_from_cdn("002", timeout=2)
    images.download_cover_from_cdn("001", timeout=3)
    assert opener.call_count == 3


@pytest.mark.parametrize(
    "error", [URLError("offline"), TimeoutError(), ValueError("invalid")]
)
def test_download_expected_failure(monkeypatch, error):
    opener = Mock(side_effect=error)
    monkeypatch.setattr(images, "urlopen", opener)
    assert images.download_cover_from_cdn("001") is None
    assert images.download_cover_from_cdn("001") is None
    opener.assert_called_once()


def test_download_unexpected_failure(monkeypatch):
    monkeypatch.setattr(images, "urlopen", Mock(side_effect=RuntimeError("unexpected")))
    with pytest.raises(RuntimeError, match="unexpected"):
        images.download_cover_from_cdn("001")


def test_send_back(slide):
    a = slide.shapes.add_textbox(0, 0, 100, 100)
    b = slide.shapes.add_textbox(0, 0, 100, 100)
    images.send_to_back(slide, b)
    assert [s.shape_id for s in slide.shapes] == [b.shape_id, a.shape_id]


@pytest.mark.parametrize("text", ["НЕТ ИЗОБРАЖЕНИЯ", "custom"])
def test_placeholder(slide, config, text):
    shape = images.add_missing_cover_placeholder(
        slide,
        config=config,
        position_left=Mm(1),
        position_top=Mm(2),
        width=Mm(30),
        height=Mm(40),
        text=text,
    )
    assert (shape.left, shape.top, shape.width, shape.height) == (
        Mm(1),
        Mm(2),
        Mm(30),
        Mm(40),
    )
    assert shape.text == text
    assert shape.fill.fore_color.rgb == config.colors.beige
    assert shape.text_frame.paragraphs[0].font.bold


def test_margin(slide, config):
    margin, icon = images.add_right_margin_icon(slide, config=config)
    assert margin.left == config.slide_size.width - config.margins.right
    assert margin.width == config.margins.right
    assert margin.height == config.slide_size.height
    assert margin.fill.fore_color.rgb == config.colors.blue
    assert icon.height == config.book_info_slide.icon_height


@pytest.mark.parametrize("available", [True, False])
def test_cover_composition(slide, config, png, monkeypatch, available):
    downloader = Mock(return_value=png() if available else None)
    monkeypatch.setattr(images, "download_cover_from_cdn", downloader)
    backdrop, cover = images.add_cover_with_backdrop(
        slide,
        config=config,
        item_code="001",
        cover_position_left=Mm(3),
        cover_position_top=Mm(4),
        cover_max_width=Mm(60),
        cover_max_height=Mm(80),
        backdrop_position_left=Mm(1),
        backdrop_position_top=Mm(2),
    )
    downloader.assert_called_once_with("001")
    assert (cover.width, cover.height) == (Mm(60), Mm(30 if available else 80))
    assert (backdrop.width, backdrop.height) == (cover.width, cover.height)
    assert (backdrop.left, backdrop.top) == (Mm(1), Mm(2))
    assert (cover.left, cover.top) == (Mm(3), Mm(4))
    assert slide.shapes[0].shape_id == backdrop.shape_id
    assert backdrop.fill.fore_color.rgb == config.colors.blue
    if not available:
        assert cover.text == "НЕТ ИЗОБРАЖЕНИЯ"
