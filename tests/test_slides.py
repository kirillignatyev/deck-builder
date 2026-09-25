from unittest.mock import Mock
from pptx.util import Pt


def test_title_orchestration(isolated_slides, monkeypatch, deck, config):
    texts = Mock()
    monkeypatch.setattr(isolated_slides, "add_safe_textbox", texts)
    slide = isolated_slides.add_title_slide(
        deck,
        config=config,
        title_text="Title",
        suptitle_text="Sub",
        attribution_text="Author",
        logo_path=config.assets.logo,
    )
    assert len(deck.slides) == 1
    assert slide.background.fill.fore_color.rgb == config.colors.blue
    assert [c.kwargs["text"] for c in texts.call_args_list] == [
        "Title",
        "Sub",
        "Author",
    ]
    assert [c.kwargs["font_size"] for c in texts.call_args_list] == [
        config.typography.header_1,
        config.typography.header_2,
        config.typography.header_3,
    ]
    assert texts.call_args_list[0].kwargs["bold"] is True


def test_book_orchestration(isolated_slides, monkeypatch, deck, config, book):
    boxes = []

    def textbox(**kwargs):
        box = Mock()
        box.text_frame.paragraphs = [Mock(), Mock()]
        boxes.append(box)
        return box

    text = Mock(side_effect=textbox)
    accent = Mock()
    cover = Mock()
    margin = Mock()
    for name, mock in [
        ("add_safe_textbox", text),
        ("add_accent_text_box", accent),
        ("add_cover_with_backdrop", cover),
        ("add_right_margin_icon", margin),
    ]:
        monkeypatch.setattr(isolated_slides, name, mock)
    slide = isolated_slides.add_bookinfo_slide(deck, config=config, book=book)
    assert len(deck.slides) == 1
    margin.assert_called_once_with(slide=slide, config=config)
    values = [c.kwargs["text"] for c in text.call_args_list]
    assert values[:4] == [book.authors, book.title, book.annotation, "Причина"]
    assert (
        values[4]
        == f"ISBN {book.isbn} • Серия: {book.series}\n60x90 • 128 с. • мягкая обложка • Красочность: 1+1 • Тираж: 1000 • Возраст: 12+"
    )
    assert [c.kwargs["text_items"] for c in accent.call_args_list] == [
        ["5 причин купить"],
        ["тег", "<не заполнен тег>"],
    ]
    for box in [boxes[2], boxes[3]]:
        assert all(p.space_before == Pt(5) for p in box.text_frame.paragraphs)
    cover.assert_called_once()
